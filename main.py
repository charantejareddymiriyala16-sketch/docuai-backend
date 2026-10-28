
import os
import io
import logging

from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from groq import Groq
from pypdf import PdfReader
from docx import Document

app = FastAPI(title="DocuAI API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

API_KEY = os.getenv("GROQ_API_KEY")
client = Groq(api_key=API_KEY) if API_KEY else None

MODEL_NAME = os.getenv("GROQ_MODEL", "llama-3.3-70b-versatile")


def extract_text(filename: str, content: bytes) -> str:
    name = filename.lower()

    try:
        if name.endswith(".pdf"):
            reader = PdfReader(io.BytesIO(content))
            return "\n".join(
                page.extract_text() or "" for page in reader.pages
            )

        if name.endswith(".docx"):
            doc = Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs]

            for table in doc.tables:
                for row in table.rows:
                    paragraphs.append(
                        " | ".join(cell.text for cell in row.cells)
                    )

            return "\n".join(paragraphs)

        if name.endswith(".txt"):
            return content.decode("utf-8-sig", errors="replace")

    except Exception:
        logging.exception("Document extraction failed")
        raise HTTPException(
            status_code=400,
            detail="Could not read this document. Check that the file is valid.",
        )

    raise HTTPException(
        status_code=400,
        detail="Unsupported file. Upload a PDF, DOCX, or TXT file.",
    )


def ask_groq(prompt: str) -> str:
    if not client:
        raise HTTPException(
            status_code=500,
            detail="Groq API key is not configured on the server.",
        )

    try:
        response = client.chat.completions.create(
            model=MODEL_NAME,
            messages=[
                {
                    "role": "user",
                    "content": prompt,
                }
            ],
            temperature=0.3,
            max_completion_tokens=2048,
        )

        answer = response.choices[0].message.content

        if not answer:
            raise HTTPException(
                status_code=502,
                detail="Groq returned no text. Please try again.",
            )

        return answer

    except HTTPException:
        raise

    except Exception as exc:
        logging.exception("Groq API request failed")
        status_code = getattr(exc, "status_code", None)

        if status_code == 429:
            raise HTTPException(
                status_code=429,
                detail="Groq rate limit reached. Please wait and try again.",
            )

        if status_code in (401, 403):
            raise HTTPException(
                status_code=502,
                detail="Groq API authentication failed. Check the server API key.",
            )

        raise HTTPException(
            status_code=502,
            detail="Groq request failed. Check the Render logs and model configuration.",
        )


@app.get("/")
def home():
    return {"message": "DocuAI API is running"}


@app.get("/health")
def health():
    return {"status": "ok"}


@app.post("/summarize")
async def summarize(file: UploadFile = File(...)):
    filename = file.filename or ""
    content = await file.read()

    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(
            status_code=413,
            detail="Maximum file size is 10 MB.",
        )

    text = extract_text(filename, content).strip()

    if not text:
        raise HTTPException(
            status_code=400,
            detail=(
                "No readable text found. Scanned PDFs need OCR, "
                "which is not included yet."
            ),
        )

    if len(text) > 100_000:
        text = text[:100_000]

    summary = ask_groq(
        "Summarize the following document clearly. Use a short overview, "
        "important points, and a conclusion. Do not invent facts.\n\n"
        f"DOCUMENT:\n{text}"
    )

    return {
        "filename": filename,
        "summary": summary,
        "document_text": text,
    }


@app.post("/chat")
async def chat(
    question: str = Form(...),
    document_text: str = Form(...),
):
    question = question.strip()
    document_text = document_text.strip()

    if not question:
        raise HTTPException(
            status_code=400,
            detail="Enter a question.",
        )

    if not document_text:
        raise HTTPException(
            status_code=400,
            detail="Document text is missing. Upload and summarize a document first.",
        )

    if len(document_text) > 100_000:
        document_text = document_text[:100_000]

    answer = ask_groq(
        "Answer the user's question using only the supplied document. "
        "If the answer is not in the document, say so clearly. "
        "Be clear and concise.\n\n"
        f"DOCUMENT:\n{document_text}\n\n"
        f"QUESTION:\n{question}"
    )

    return {"answer": answer}
